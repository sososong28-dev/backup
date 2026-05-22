from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import re
import shutil
import sqlite3
import sys
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable

import xlrd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SCHEMA_PATH = ROOT / "schema.sql"
CONFIG_PATH = ROOT / "config" / "config.json"
CONFIG_EXAMPLE_PATH = ROOT / "config" / "config.example.json"


def now_iso() -> str:
    return dt.datetime.now().replace(microsecond=0).isoformat(sep=" ")


def today() -> str:
    return dt.date.today().isoformat()


def canonical(value: str) -> str:
    return re.sub(r"\s+", "", str(value).strip()).lower()


def clean_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.replace(microsecond=0).isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    return value


def clean_text(value: Any) -> str | None:
    value = clean_cell(value)
    if value is None or value == "":
        return None
    return str(value).strip()


def clean_number(value: Any) -> float | None:
    value = clean_cell(value)
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    text = text.replace(",", "")
    text = re.sub(r"[￥¥$元件个\s]", "", text)
    try:
        return float(text)
    except ValueError:
        return None


def load_config(path: Path | None = None) -> dict[str, Any]:
    config_path = path or CONFIG_PATH
    if not config_path.exists():
        config_path = CONFIG_EXAMPLE_PATH
    with config_path.open("r", encoding="utf-8") as f:
        config = json.load(f)

    for key in ("database_path", "inbox_dir", "archive_dir", "error_dir"):
        value = Path(config[key])
        if not value.is_absolute():
            value = ROOT / value
        config[key] = str(value)
    return config


def connect_db(config: dict[str, Any]) -> sqlite3.Connection:
    db_path = Path(config["database_path"])
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db(conn: sqlite3.Connection) -> None:
    with SCHEMA_PATH.open("r", encoding="utf-8") as f:
        conn.executescript(f.read())
    conn.commit()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def row_hash(*parts: Any) -> str:
    text = "|".join("" if part is None else str(part) for part in parts)
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def unique_headers(values: Iterable[Any]) -> list[str]:
    counts: dict[str, int] = {}
    headers: list[str] = []
    for idx, value in enumerate(values, start=1):
        header = clean_text(value) or f"column_{idx}"
        base = header
        counts[base] = counts.get(base, 0) + 1
        if counts[base] > 1:
            header = f"{base}_{counts[base]}"
        headers.append(header)
    return headers


def non_empty_count(values: Iterable[Any]) -> int:
    return sum(1 for value in values if clean_text(value))


def detect_header_row(rows: list[list[Any]]) -> int:
    best_idx = 0
    best_count = 0
    for idx, row in enumerate(rows[:20]):
        count = non_empty_count(row)
        if count > best_count:
            best_idx = idx
            best_count = count
    return best_idx


def trim_row(row: list[Any], width: int) -> list[Any]:
    padded = list(row[:width])
    if len(padded) < width:
        padded.extend([""] * (width - len(padded)))
    return [clean_cell(value) for value in padded]


def read_xlsx(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        raw_rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        raw_rows = [row for row in raw_rows if non_empty_count(row) > 0]
        if not raw_rows:
            continue
        header_idx = detect_header_row(raw_rows)
        headers = unique_headers(raw_rows[header_idx])
        data_rows = []
        for row_number, row in enumerate(raw_rows[header_idx + 1 :], start=header_idx + 2):
            if non_empty_count(row) == 0:
                continue
            values = trim_row(row, len(headers))
            data_rows.append((row_number, dict(zip(headers, values))))
        sheets.append(
            {
                "sheet_name": worksheet.title,
                "headers": headers,
                "rows": data_rows,
            }
        )
    workbook.close()
    return sheets


class HtmlTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.rows: list[list[str]] = []
        self._in_cell = False
        self._cell_parts: list[str] = []
        self._row: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        tag = tag.lower()
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._in_cell = True
            self._cell_parts = []

    def handle_data(self, data: str) -> None:
        if self._in_cell:
            self._cell_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in ("td", "th") and self._row is not None and self._in_cell:
            self._row.append("".join(self._cell_parts).strip())
            self._in_cell = False
            self._cell_parts = []
        elif tag == "tr" and self._row is not None:
            if non_empty_count(self._row) > 0:
                self.rows.append(self._row)
            self._row = None


def rows_to_sheet(rows: list[list[Any]], sheet_name: str = "") -> list[dict[str, Any]]:
    rows = [row for row in rows if non_empty_count(row) > 0]
    if not rows:
        return []
    header_idx = detect_header_row(rows)
    headers = unique_headers(rows[header_idx])
    data_rows = []
    for row_number, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if non_empty_count(row) == 0:
            continue
        values = trim_row(row, len(headers))
        data_rows.append((row_number, dict(zip(headers, values))))
    return [{"sheet_name": sheet_name, "headers": headers, "rows": data_rows}]


def read_xls(path: Path, preferred_encoding: str) -> list[dict[str, Any]]:
    prefix = path.read_bytes()[:8]
    if prefix.startswith(b"\xd0\xcf"):
        workbook = xlrd.open_workbook(str(path))
        sheets: list[dict[str, Any]] = []
        for sheet in workbook.sheets():
            raw_rows = [sheet.row_values(idx) for idx in range(sheet.nrows)]
            sheets.extend(rows_to_sheet(raw_rows, sheet.name))
        return sheets

    encoding = detect_csv_encoding(path, preferred_encoding)
    text = path.read_text(encoding=encoding)
    if "<table" in text.lower():
        parser = HtmlTableParser()
        parser.feed(text)
        return rows_to_sheet(parser.rows)

    return read_csv(path, preferred_encoding)


def detect_csv_encoding(path: Path, preferred: str) -> str:
    for encoding in (preferred, "utf-8-sig", "gb18030"):
        try:
            path.read_text(encoding=encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    return preferred


def read_csv(path: Path, preferred_encoding: str) -> list[dict[str, Any]]:
    encoding = detect_csv_encoding(path, preferred_encoding)
    text = path.read_text(encoding=encoding)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample)
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(text.splitlines(), dialect)
    rows = [row for row in reader if non_empty_count(row) > 0]
    return rows_to_sheet(rows)


def read_export(path: Path, config: dict[str, Any]) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return read_xlsx(path)
    if suffix == ".xls":
        return read_xls(path, config.get("default_encoding", "utf-8-sig"))
    if suffix == ".csv":
        return read_csv(path, config.get("default_encoding", "utf-8-sig"))
    raise ValueError(f"Unsupported export file type: {path.suffix}")


def get_mapping(config: dict[str, Any], entity_type: str) -> dict[str, list[str]]:
    return config["export_types"].get(entity_type, {}).get("mapping", {})


def value_by_alias(row: dict[str, Any], aliases: list[str] | None) -> Any:
    if not aliases:
        return None
    lookup = {canonical(key): value for key, value in row.items()}
    for alias in aliases:
        key = canonical(alias)
        if key in lookup and clean_text(lookup[key]) is not None:
            return lookup[key]
    return None


def mapped(row: dict[str, Any], mapping: dict[str, list[str]], field: str) -> Any:
    return value_by_alias(row, mapping.get(field, []))


def mapped_text(row: dict[str, Any], mapping: dict[str, list[str]], field: str) -> str | None:
    return clean_text(mapped(row, mapping, field))


def mapped_number(row: dict[str, Any], mapping: dict[str, list[str]], field: str) -> float | None:
    return clean_number(mapped(row, mapping, field))


def first_present(*values: Any) -> str | None:
    for value in values:
        text = clean_text(value)
        if text:
            return text
    return None


def raw_json(row: dict[str, Any]) -> str:
    return json.dumps(row, ensure_ascii=False, default=str, sort_keys=True)


def detect_entity(
    path: Path, headers: list[str], config: dict[str, Any], forced: str | None = None
) -> str:
    if forced:
        return forced

    name = path.name.lower()
    best_entity = "unknown"
    best_score = 0
    for entity, spec in config["export_types"].items():
        for pattern in spec.get("file_patterns", []):
            if pattern.lower() in name:
                return entity

        header_keys = {canonical(header) for header in headers}
        aliases = []
        for field_aliases in spec.get("mapping", {}).values():
            aliases.extend(field_aliases)
        score = sum(1 for alias in aliases if canonical(alias) in header_keys)
        if score > best_score:
            best_score = score
            best_entity = entity

    return best_entity if best_score > 0 else "unknown"


def upsert_shop(
    conn: sqlite3.Connection,
    row: dict[str, Any],
    mapping: dict[str, list[str]],
    updated_at: str,
) -> str | None:
    shop_name = mapped_text(row, mapping, "shop_name")
    shop_no = mapped_text(row, mapping, "shop_no")
    shop_key = first_present(shop_no, shop_name)
    if not shop_key:
        return None
    conn.execute(
        """
        INSERT INTO dim_shop (
          shop_key, shop_no, shop_name, platform_name, raw_json, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?)
        ON CONFLICT(shop_key) DO UPDATE SET
          shop_no = excluded.shop_no,
          shop_name = excluded.shop_name,
          platform_name = excluded.platform_name,
          raw_json = excluded.raw_json,
          updated_at = excluded.updated_at
        """,
        (
            shop_key,
            shop_no,
            shop_name or shop_key,
            mapped_text(row, mapping, "platform_name"),
            raw_json(row),
            updated_at,
        ),
    )
    return shop_key


def upsert_warehouse(
    conn: sqlite3.Connection,
    row: dict[str, Any],
    mapping: dict[str, list[str]],
    updated_at: str,
) -> str | None:
    warehouse_name = mapped_text(row, mapping, "warehouse_name")
    warehouse_no = mapped_text(row, mapping, "warehouse_no")
    warehouse_key = first_present(warehouse_no, warehouse_name)
    if not warehouse_key:
        return None
    conn.execute(
        """
        INSERT INTO dim_warehouse (
          warehouse_key, warehouse_no, warehouse_name, raw_json, updated_at
        )
        VALUES (?, ?, ?, ?, ?)
        ON CONFLICT(warehouse_key) DO UPDATE SET
          warehouse_no = excluded.warehouse_no,
          warehouse_name = excluded.warehouse_name,
          raw_json = excluded.raw_json,
          updated_at = excluded.updated_at
        """,
        (
            warehouse_key,
            warehouse_no,
            warehouse_name or warehouse_key,
            raw_json(row),
            updated_at,
        ),
    )
    return warehouse_key


def upsert_goods(
    conn: sqlite3.Connection,
    row: dict[str, Any],
    mapping: dict[str, list[str]],
    updated_at: str,
    fallback: str,
) -> str:
    sku_key = first_present(
        mapped_text(row, mapping, "sku_key"),
        mapped_text(row, mapping, "spec_no"),
        mapped_text(row, mapping, "goods_no"),
        mapped_text(row, mapping, "barcode"),
        fallback,
    )
    assert sku_key is not None
    conn.execute(
        """
        INSERT INTO dim_goods (
          sku_key, goods_no, spec_no, barcode, goods_name, spec_name,
          category_name, brand_name, raw_json, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(sku_key) DO UPDATE SET
          goods_no = coalesce(excluded.goods_no, dim_goods.goods_no),
          spec_no = coalesce(excluded.spec_no, dim_goods.spec_no),
          barcode = coalesce(excluded.barcode, dim_goods.barcode),
          goods_name = coalesce(excluded.goods_name, dim_goods.goods_name),
          spec_name = coalesce(excluded.spec_name, dim_goods.spec_name),
          category_name = coalesce(excluded.category_name, dim_goods.category_name),
          brand_name = coalesce(excluded.brand_name, dim_goods.brand_name),
          raw_json = excluded.raw_json,
          updated_at = excluded.updated_at
        """,
        (
            sku_key,
            mapped_text(row, mapping, "goods_no"),
            mapped_text(row, mapping, "spec_no"),
            mapped_text(row, mapping, "barcode"),
            mapped_text(row, mapping, "goods_name"),
            mapped_text(row, mapping, "spec_name"),
            mapped_text(row, mapping, "category_name"),
            mapped_text(row, mapping, "brand_name"),
            raw_json(row),
            updated_at,
        ),
    )
    return sku_key


def upsert_trade_stub(
    conn: sqlite3.Connection,
    trade_no: str,
    source_file_id: int,
    updated_at: str,
) -> None:
    conn.execute(
        """
        INSERT INTO fact_trade (trade_no, source_file_id, updated_at)
        VALUES (?, ?, ?)
        ON CONFLICT(trade_no) DO NOTHING
        """,
        (trade_no, source_file_id, updated_at),
    )


def import_shop_row(
    conn: sqlite3.Connection,
    row: dict[str, Any],
    source_file_id: int,
    source_hash: str,
    config: dict[str, Any],
    updated_at: str,
) -> None:
    del source_file_id, source_hash
    mapping = get_mapping(config, "shops")
    upsert_shop(conn, row, mapping, updated_at)


def import_goods_row(
    conn: sqlite3.Connection,
    row: dict[str, Any],
    source_file_id: int,
    source_hash: str,
    config: dict[str, Any],
    updated_at: str,
) -> None:
    del source_file_id
    mapping = get_mapping(config, "goods")
    upsert_goods(conn, row, mapping, updated_at, source_hash)


def import_order_row(
    conn: sqlite3.Connection,
    row: dict[str, Any],
    source_file_id: int,
    source_hash: str,
    config: dict[str, Any],
    updated_at: str,
) -> None:
    mapping = get_mapping(config, "orders")
    trade_no = first_present(mapped_text(row, mapping, "trade_no"), source_hash)
    assert trade_no is not None
    shop_key = upsert_shop(conn, row, mapping, updated_at)

    conn.execute(
        """
        INSERT INTO fact_trade (
          trade_no, src_tid, shop_key, trade_status, trade_time, pay_time,
          buyer_nick, receiver_name, receiver_mobile_masked, goods_amount,
          post_amount, paid_amount, source_file_id, raw_json, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(trade_no) DO UPDATE SET
          src_tid = coalesce(excluded.src_tid, fact_trade.src_tid),
          shop_key = coalesce(excluded.shop_key, fact_trade.shop_key),
          trade_status = coalesce(excluded.trade_status, fact_trade.trade_status),
          trade_time = coalesce(excluded.trade_time, fact_trade.trade_time),
          pay_time = coalesce(excluded.pay_time, fact_trade.pay_time),
          buyer_nick = coalesce(excluded.buyer_nick, fact_trade.buyer_nick),
          receiver_name = coalesce(excluded.receiver_name, fact_trade.receiver_name),
          receiver_mobile_masked = coalesce(
            excluded.receiver_mobile_masked, fact_trade.receiver_mobile_masked
          ),
          goods_amount = coalesce(excluded.goods_amount, fact_trade.goods_amount),
          post_amount = coalesce(excluded.post_amount, fact_trade.post_amount),
          paid_amount = coalesce(excluded.paid_amount, fact_trade.paid_amount),
          source_file_id = excluded.source_file_id,
          raw_json = excluded.raw_json,
          updated_at = excluded.updated_at
        """,
        (
            trade_no,
            mapped_text(row, mapping, "src_tid"),
            shop_key,
            mapped_text(row, mapping, "trade_status"),
            mapped_text(row, mapping, "trade_time"),
            mapped_text(row, mapping, "pay_time"),
            mapped_text(row, mapping, "buyer_nick"),
            mapped_text(row, mapping, "receiver_name"),
            mapped_text(row, mapping, "receiver_mobile_masked"),
            mapped_number(row, mapping, "goods_amount"),
            mapped_number(row, mapping, "post_amount"),
            mapped_number(row, mapping, "paid_amount"),
            source_file_id,
            raw_json(row),
            updated_at,
        ),
    )

    if any(mapped_text(row, mapping, field) for field in ("sku_key", "goods_no", "goods_name")):
        sku_key = upsert_goods(conn, row, mapping, updated_at, source_hash)
        item_id = first_present(mapped_text(row, mapping, "item_id"), source_hash)
        item_key = row_hash(trade_no, sku_key, item_id)
        conn.execute(
            """
            INSERT INTO fact_trade_item (
              item_key, trade_no, sku_key, goods_count, price, discount_amount,
              paid_amount, source_file_id, raw_json, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(item_key) DO UPDATE SET
              goods_count = coalesce(excluded.goods_count, fact_trade_item.goods_count),
              price = coalesce(excluded.price, fact_trade_item.price),
              discount_amount = coalesce(
                excluded.discount_amount, fact_trade_item.discount_amount
              ),
              paid_amount = coalesce(excluded.paid_amount, fact_trade_item.paid_amount),
              source_file_id = excluded.source_file_id,
              raw_json = excluded.raw_json,
              updated_at = excluded.updated_at
            """,
            (
                item_key,
                trade_no,
                sku_key,
                mapped_number(row, mapping, "goods_count"),
                mapped_number(row, mapping, "price"),
                mapped_number(row, mapping, "discount_amount"),
                mapped_number(row, mapping, "item_paid_amount"),
                source_file_id,
                raw_json(row),
                updated_at,
            ),
        )


def import_order_item_row(
    conn: sqlite3.Connection,
    row: dict[str, Any],
    source_file_id: int,
    source_hash: str,
    config: dict[str, Any],
    updated_at: str,
) -> None:
    mapping = get_mapping(config, "order_items")
    trade_no = mapped_text(row, mapping, "trade_no")
    if trade_no:
        upsert_trade_stub(conn, trade_no, source_file_id, updated_at)
    sku_key = upsert_goods(conn, row, mapping, updated_at, source_hash)
    item_id = first_present(mapped_text(row, mapping, "item_id"), source_hash)
    item_key = row_hash(trade_no, sku_key, item_id)
    conn.execute(
        """
        INSERT INTO fact_trade_item (
          item_key, trade_no, sku_key, goods_count, price, discount_amount,
          paid_amount, source_file_id, raw_json, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(item_key) DO UPDATE SET
          goods_count = coalesce(excluded.goods_count, fact_trade_item.goods_count),
          price = coalesce(excluded.price, fact_trade_item.price),
          discount_amount = coalesce(
            excluded.discount_amount, fact_trade_item.discount_amount
          ),
          paid_amount = coalesce(excluded.paid_amount, fact_trade_item.paid_amount),
          source_file_id = excluded.source_file_id,
          raw_json = excluded.raw_json,
          updated_at = excluded.updated_at
        """,
        (
            item_key,
            trade_no,
            sku_key,
            mapped_number(row, mapping, "goods_count"),
            mapped_number(row, mapping, "price"),
            mapped_number(row, mapping, "discount_amount"),
            mapped_number(row, mapping, "paid_amount"),
            source_file_id,
            raw_json(row),
            updated_at,
        ),
    )


def import_stock_row(
    conn: sqlite3.Connection,
    row: dict[str, Any],
    source_file_id: int,
    source_hash: str,
    config: dict[str, Any],
    updated_at: str,
) -> None:
    mapping = get_mapping(config, "stock")
    warehouse_key = upsert_warehouse(conn, row, mapping, updated_at)
    sku_key = upsert_goods(conn, row, mapping, updated_at, source_hash)
    snapshot_date = mapped_text(row, mapping, "snapshot_date") or today()
    snapshot_date = snapshot_date[:10]
    snapshot_key = row_hash(snapshot_date, warehouse_key, sku_key)
    conn.execute(
        """
        INSERT INTO fact_stock_snapshot (
          snapshot_key, snapshot_date, warehouse_key, sku_key, stock_qty,
          available_qty, occupy_qty, source_file_id, raw_json, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(snapshot_key) DO UPDATE SET
          stock_qty = coalesce(excluded.stock_qty, fact_stock_snapshot.stock_qty),
          available_qty = coalesce(
            excluded.available_qty, fact_stock_snapshot.available_qty
          ),
          occupy_qty = coalesce(excluded.occupy_qty, fact_stock_snapshot.occupy_qty),
          source_file_id = excluded.source_file_id,
          raw_json = excluded.raw_json,
          updated_at = excluded.updated_at
        """,
        (
            snapshot_key,
            snapshot_date,
            warehouse_key,
            sku_key,
            mapped_number(row, mapping, "stock_qty"),
            mapped_number(row, mapping, "available_qty"),
            mapped_number(row, mapping, "occupy_qty"),
            source_file_id,
            raw_json(row),
            updated_at,
        ),
    )


def import_refund_row(
    conn: sqlite3.Connection,
    row: dict[str, Any],
    source_file_id: int,
    source_hash: str,
    config: dict[str, Any],
    updated_at: str,
) -> None:
    mapping = get_mapping(config, "refunds")
    refund_no = first_present(mapped_text(row, mapping, "refund_no"), source_hash)
    assert refund_no is not None
    shop_key = upsert_shop(conn, row, mapping, updated_at)
    conn.execute(
        """
        INSERT INTO fact_refund (
          refund_no, trade_no, shop_key, refund_status, refund_time,
          refund_amount, reason, source_file_id, raw_json, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(refund_no) DO UPDATE SET
          trade_no = coalesce(excluded.trade_no, fact_refund.trade_no),
          shop_key = coalesce(excluded.shop_key, fact_refund.shop_key),
          refund_status = coalesce(excluded.refund_status, fact_refund.refund_status),
          refund_time = coalesce(excluded.refund_time, fact_refund.refund_time),
          refund_amount = coalesce(excluded.refund_amount, fact_refund.refund_amount),
          reason = coalesce(excluded.reason, fact_refund.reason),
          source_file_id = excluded.source_file_id,
          raw_json = excluded.raw_json,
          updated_at = excluded.updated_at
        """,
        (
            refund_no,
            mapped_text(row, mapping, "trade_no"),
            shop_key,
            mapped_text(row, mapping, "refund_status"),
            mapped_text(row, mapping, "refund_time"),
            mapped_number(row, mapping, "refund_amount"),
            mapped_text(row, mapping, "reason"),
            source_file_id,
            raw_json(row),
            updated_at,
        ),
    )


IMPORTERS = {
    "goods": import_goods_row,
    "orders": import_order_row,
    "order_items": import_order_item_row,
    "stock": import_stock_row,
    "refunds": import_refund_row,
    "shops": import_shop_row,
}


def insert_raw_row(
    conn: sqlite3.Connection,
    source_file_id: int,
    entity_type: str,
    sheet_name: str,
    row_number: int,
    source_hash: str,
    row: dict[str, Any],
    ingested_at: str,
) -> None:
    conn.execute(
        """
        INSERT OR IGNORE INTO raw_wdt_rows (
          source_file_id, entity_type, sheet_name, row_number, source_row_hash,
          row_json, ingested_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (
            source_file_id,
            entity_type,
            sheet_name,
            row_number,
            source_hash,
            raw_json(row),
            ingested_at,
        ),
    )


def file_record(
    conn: sqlite3.Connection, file_hash: str, sheet_name: str
) -> sqlite3.Row | None:
    return conn.execute(
        """
        SELECT *
        FROM rpa_export_files
        WHERE file_sha256 = ? AND sheet_name = ?
        """,
        (file_hash, sheet_name),
    ).fetchone()


def reuse_file_record(
    conn: sqlite3.Connection,
    source_file_id: int,
    run_id: int,
    path: Path,
    entity_type: str,
    row_count: int = 0,
    status: str = "importing",
    error_message: str | None = None,
) -> None:
    conn.execute(
        """
        UPDATE rpa_export_files
        SET run_id = ?, file_path = ?, entity_type = ?, status = ?,
            row_count = ?, imported_at = ?, error_message = ?
        WHERE id = ?
        """,
        (
            run_id,
            str(path),
            entity_type,
            status,
            row_count,
            now_iso(),
            error_message,
            source_file_id,
        ),
    )


def insert_file_record(
    conn: sqlite3.Connection,
    run_id: int,
    path: Path,
    file_hash: str,
    entity_type: str,
    sheet_name: str,
    status: str,
    row_count: int = 0,
    error_message: str | None = None,
) -> int:
    stat = path.stat()
    conn.execute(
        """
        INSERT INTO rpa_export_files (
          run_id, file_path, original_file_name, file_sha256, file_size,
          file_mtime, entity_type, sheet_name, status, row_count, imported_at,
          error_message
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            run_id,
            str(path),
            path.name,
            file_hash,
            stat.st_size,
            dt.datetime.fromtimestamp(stat.st_mtime).replace(microsecond=0).isoformat(sep=" "),
            entity_type,
            sheet_name,
            status,
            row_count,
            now_iso(),
            error_message,
        ),
    )
    return int(conn.execute("SELECT last_insert_rowid()").fetchone()[0])


def import_one_file(
    conn: sqlite3.Connection,
    run_id: int,
    path: Path,
    config: dict[str, Any],
    force: bool = False,
    forced_entity: str | None = None,
) -> tuple[int, int, int]:
    file_hash = sha256_file(path)
    imported_files = 0
    imported_rows = 0
    skipped_files = 0
    sheets = read_export(path, config)
    if not sheets:
        raise ValueError("No rows found")

    for sheet in sheets:
        sheet_name = sheet["sheet_name"] or ""
        headers = sheet["headers"]
        rows = sheet["rows"]
        entity_type = detect_entity(path, headers, config, forced_entity)

        prior = file_record(conn, file_hash, sheet_name)
        if prior and prior["status"] == "success" and not force:
            skipped_files += 1
            continue
        if prior:
            source_file_id = int(prior["id"])
            reuse_file_record(conn, source_file_id, run_id, path, entity_type)
        else:
            source_file_id = insert_file_record(
                conn,
                run_id,
                path,
                file_hash,
                entity_type,
                sheet_name,
                "importing",
            )
        importer = IMPORTERS.get(entity_type)
        ingested_at = now_iso()

        for row_number, row in rows:
            source_hash = row_hash(file_hash, sheet_name, row_number, raw_json(row))
            insert_raw_row(
                conn,
                source_file_id,
                entity_type,
                sheet_name,
                row_number,
                source_hash,
                row,
                ingested_at,
            )
            if importer:
                importer(conn, row, source_file_id, source_hash, config, ingested_at)
            imported_rows += 1

        conn.execute(
            """
            UPDATE rpa_export_files
            SET status = 'success', row_count = ?, imported_at = ?
            WHERE id = ?
            """,
            (len(rows), now_iso(), source_file_id),
        )
        imported_files += 1

    return imported_files, imported_rows, skipped_files


def archive_file(path: Path, config: dict[str, Any], ok: bool) -> None:
    if not config.get("archive_after_import", True):
        return
    target_root = Path(config["archive_dir"] if ok else config["error_dir"])
    target_root.mkdir(parents=True, exist_ok=True)
    target = target_root / path.name
    if target.exists():
        stem = path.stem
        suffix = path.suffix
        target = target_root / f"{stem}_{dt.datetime.now().strftime('%Y%m%d_%H%M%S')}{suffix}"
    shutil.move(str(path), str(target))


def import_inbox(
    config: dict[str, Any],
    force: bool = False,
    forced_entity: str | None = None,
) -> int:
    inbox = Path(config["inbox_dir"])
    inbox.mkdir(parents=True, exist_ok=True)

    conn = connect_db(config)
    init_db(conn)

    run_started = now_iso()
    conn.execute(
        """
        INSERT INTO sync_runs (run_started_at, status, inbox_path)
        VALUES (?, 'running', ?)
        """,
        (run_started, str(inbox)),
    )
    run_id = int(conn.execute("SELECT last_insert_rowid()").fetchone()[0])
    conn.commit()

    imported_files = 0
    imported_rows = 0
    skipped_files = 0
    errors: list[str] = []

    patterns = ("*.csv", "*.xlsx", "*.xlsm", "*.xls")
    files = sorted(path for pattern in patterns for path in inbox.glob(pattern))
    for path in files:
        ok = False
        try:
            with conn:
                f_count, r_count, s_count = import_one_file(
                    conn, run_id, path, config, force=force, forced_entity=forced_entity
                )
            imported_files += f_count
            imported_rows += r_count
            skipped_files += s_count
            ok = True
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{path.name}: {exc}")
            try:
                file_hash = sha256_file(path)
                prior = file_record(conn, file_hash, "")
                if prior:
                    reuse_file_record(
                        conn,
                        int(prior["id"]),
                        run_id,
                        path,
                        forced_entity or prior["entity_type"] or "unknown",
                        status="error",
                        error_message=str(exc),
                    )
                else:
                    insert_file_record(
                        conn,
                        run_id,
                        path,
                        file_hash,
                        forced_entity or "unknown",
                        "",
                        "error",
                        error_message=str(exc),
                    )
                conn.commit()
            except Exception:
                conn.rollback()
        finally:
            if ok or errors:
                archive_file(path, config, ok)

    status = "success" if not errors else "error"
    conn.execute(
        """
        UPDATE sync_runs
        SET run_finished_at = ?, status = ?, imported_files = ?,
            imported_rows = ?, skipped_files = ?, error_message = ?
        WHERE id = ?
        """,
        (
            now_iso(),
            status,
            imported_files,
            imported_rows,
            skipped_files,
            "\n".join(errors) if errors else None,
            run_id,
        ),
    )
    conn.commit()
    conn.close()

    print(
        json.dumps(
            {
                "status": status,
                "run_id": run_id,
                "imported_files": imported_files,
                "imported_rows": imported_rows,
                "skipped_files": skipped_files,
                "errors": errors,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0 if not errors else 1


def main() -> int:
    parser = argparse.ArgumentParser(description="Import WDT RPA exports into SQLite.")
    parser.add_argument("--config", type=Path, default=None)
    parser.add_argument("--force", action="store_true", help="Import duplicate files again.")
    parser.add_argument(
        "--entity",
        choices=sorted(IMPORTERS.keys()),
        help="Force one entity type for all files in inbox.",
    )
    args = parser.parse_args()

    config = load_config(args.config)
    return import_inbox(config, force=args.force, forced_entity=args.entity)


if __name__ == "__main__":
    sys.exit(main())
